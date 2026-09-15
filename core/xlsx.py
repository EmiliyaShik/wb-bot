"""Книги xlsx: чтение и запись листов с заголовками.

Общий помощник, а не часть какого-то одного отчёта. Здесь нет ни
себестоимости, ни финансовых полей, ни реестра актов: модуль знает только
про книгу, лист, строку заголовков и строки данных. Всё, что зависит от
предметной области, живёт у того, кто зовёт.

Две стороны:

* запись - `write_book(sheets, path=None)`. Заголовок жирный и закреплён,
  ширина колонок считается по содержимому. Без `path` возвращает байты,
  готовые к отправке в Telegram, с `path` ещё и кладёт файл на диск;
* чтение - `read_book(source)` и `read_sheet(source)`. Приходит книга из
  листов, у листа заголовки и строки, у строки настоящий номер, как его
  видит человек в Excel. Номер нужен, чтобы сказать «строка 12: ...».

Приём файла от постороннего это поверхность для атаки, поэтому у чтения
три предела: размер архива до разбора, формат до открытия и число строк и
ячеек уже внутри. Последнее важнее, чем кажется: короткий zip разжимается
в лист на миллионы строк. Чужая начинка, переименованная в `.xlsx`, даёт
понятную ошибку, а не падение openpyxl, а настоящая причина беды уходит в
журнал владельца.
"""

from __future__ import annotations

import io
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any, Iterator, Mapping, Sequence

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from core import audit

__all__ = [
    "XlsxError",
    "NotXlsxError",
    "TooLargeError",
    "TooManyRowsError",
    "TooManyCellsError",
    "SheetNotFoundError",
    "Sheet",
    "Row",
    "SheetData",
    "Book",
    "write_book",
    "read_book",
    "read_sheet",
    "looks_like_xlsx",
    "normalize",
    "find_column",
    "HEADER_ROW",
    "FIRST_DATA_ROW",
    "MIN_WIDTH",
    "MAX_WIDTH",
    "MAX_ROWS",
    "MAX_CELLS",
]

# Строка заголовков в листе одна и всегда первая: так книгу читает человек,
# так её собирает Excel по умолчанию, и так не нужен параметр «а с какой
# строки начинать».
HEADER_ROW = 1
FIRST_DATA_ROW = HEADER_ROW + 1

MIN_WIDTH = 9
MAX_WIDTH = 60

# Потолок на разжатое содержимое. Размер архива тут не помощник: короткий
# zip разворачивается в лист на миллионы строк, и книга целиком уезжает в
# память. Значения с запасом для любого разумного прайса и шаблона.
MAX_ROWS = 100_000
MAX_CELLS = 2_000_000


class XlsxError(Exception):
    """Общий предок ошибок этого модуля: одного except хватает на все."""


class NotXlsxError(XlsxError):
    """Это не книга xlsx, как бы ни назывался файл."""


class TooLargeError(XlsxError):
    """Файл больше разрешённого размера. Поднимается до разбора."""

    def __init__(self, size: int, limit: int) -> None:
        super().__init__(f"файл {size} байт при лимите {limit}")
        self.size = size
        self.limit = limit


class TooManyRowsError(XlsxError):
    """В книге больше строк, чем разрешено держать в памяти."""

    def __init__(self, rows: int, limit: int) -> None:
        super().__init__(f"строк больше {limit}")
        self.rows = rows
        self.limit = limit


class TooManyCellsError(XlsxError):
    """В книге больше ячеек, чем разрешено держать в памяти."""

    def __init__(self, cells: int, limit: int) -> None:
        super().__init__(f"ячеек больше {limit}")
        self.cells = cells
        self.limit = limit


class SheetNotFoundError(XlsxError):
    """В книге нет листа с таким названием."""


@dataclass(frozen=True)
class Sheet:
    """Лист на запись: название, заголовки, строки значений.

    `widths` задаёт ширину колонок вручную, если считать по содержимому не
    хочется. Длина списка может быть короче, чем заголовков: остальные
    посчитаются сами.
    """

    title: str
    headers: Sequence[str]
    rows: Sequence[Sequence[Any]] = ()
    widths: Sequence[int] | None = None


@dataclass(frozen=True)
class Row:
    """Строка данных: настоящий номер в листе и значения по заголовкам."""

    number: int
    cells: Mapping[str, Any]
    values: tuple[Any, ...] = ()

    def get(self, header: str, default: Any = None) -> Any:
        return self.cells.get(header, default)


@dataclass(frozen=True)
class SheetData:
    """Прочитанный лист."""

    title: str
    headers: tuple[str, ...]
    rows: tuple[Row, ...] = ()

    def column(self, *aliases: str) -> str | None:
        """Заголовок по любому из синонимов, без учёта регистра и пробелов."""
        return find_column(self.headers, aliases)

    def missing(self, required: Mapping[str, Sequence[str]]) -> tuple[str, ...]:
        """Какие из нужных колонок не нашлись. Ключ это имя для человека."""
        return tuple(
            name
            for name, aliases in required.items()
            if find_column(self.headers, aliases) is None
        )


@dataclass(frozen=True)
class Book:
    """Прочитанная книга: листы в том порядке, в каком они лежат в файле."""

    sheets: tuple[SheetData, ...] = field(default_factory=tuple)

    @property
    def titles(self) -> tuple[str, ...]:
        return tuple(sheet.title for sheet in self.sheets)

    @property
    def first(self) -> SheetData:
        if not self.sheets:
            raise SheetNotFoundError("в книге нет ни одного листа")
        return self.sheets[0]

    def get(self, title: str) -> SheetData | None:
        wanted = normalize(title)
        for sheet in self.sheets:
            if normalize(sheet.title) == wanted:
                return sheet
        return None

    def __getitem__(self, title: str) -> SheetData:
        sheet = self.get(title)
        if sheet is None:
            raise SheetNotFoundError(f"в книге нет листа «{title}»")
        return sheet

    def __iter__(self) -> Iterator[SheetData]:
        return iter(self.sheets)

    def __len__(self) -> int:
        return len(self.sheets)


def normalize(text: Any) -> str:
    """Заголовок к сравнимому виду: без регистра, без краёв, один пробел."""
    return " ".join(str(text or "").split()).casefold()


def find_column(headers: Sequence[str], aliases: Sequence[str]) -> str | None:
    """Первый заголовок, совпавший с любым из синонимов.

    Возвращается заголовок в том виде, в каком он в файле: им потом
    доставать значение из `Row.cells`.
    """
    wanted = [normalize(alias) for alias in aliases]
    for header in headers:
        if normalize(header) in wanted:
            return header
    return None


# --- запись ---


def write_book(
    sheets: Sheet | Sequence[Sheet], *, path: str | Path | None = None
) -> bytes:
    """Собирает книгу и возвращает её байтами.

    Байты, а не открытый файл: книгу чаще отправляют, чем сохраняют. Если
    нужен и файл на диске, дайте `path`, байты вернутся всё равно.
    """
    wanted = [sheets] if isinstance(sheets, Sheet) else list(sheets)
    if not wanted:
        raise XlsxError("книга без листов не бывает")

    book = Workbook()
    book.remove(book.active)
    for sheet in wanted:
        _fill(book.create_sheet(_safe_title(sheet.title)), sheet)

    buffer = io.BytesIO()
    book.save(buffer)
    data = buffer.getvalue()
    if path is not None:
        Path(path).write_bytes(data)
    return data


# Excel запрещает эти символы в названии листа и обрезает его до 31 знака.
_BAD_TITLE_CHARS = set(r"[]:*?/\\")


def _safe_title(title: str) -> str:
    cleaned = "".join(" " if ch in _BAD_TITLE_CHARS else ch for ch in str(title))
    return cleaned.strip()[:31] or "Лист"


def _fill(worksheet: Any, sheet: Sheet) -> None:
    headers = [str(header) for header in sheet.headers]
    worksheet.append(headers)
    for cell in worksheet[HEADER_ROW]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(vertical="center", wrap_text=True)

    lengths = [len(header) for header in headers]
    for row in sheet.rows:
        values = list(row)
        worksheet.append(values)
        for index, value in enumerate(values):
            if index >= len(lengths):
                lengths.append(0)
            lengths[index] = max(lengths[index], len(str(value if value is not None else "")))

    given = list(sheet.widths or ())
    for index in range(len(lengths)):
        if index < len(given) and given[index]:
            width = int(given[index])
        else:
            width = min(MAX_WIDTH, max(MIN_WIDTH, lengths[index] + 2))
        worksheet.column_dimensions[get_column_letter(index + 1)].width = width

    if headers:
        worksheet.freeze_panes = f"A{FIRST_DATA_ROW}"


# --- чтение ---


def looks_like_xlsx(data: bytes) -> bool:
    """Правда ли внутри книга, а не чужой формат с расширением .xlsx.

    Смотрим не на имя файла, а на содержимое: zip-архив, внутри которого
    лежит `xl/workbook.xml`. Переименованный pdf или csv сюда не пройдёт.
    """
    if not isinstance(data, (bytes, bytearray)) or data[:2] != b"PK":
        return False
    try:
        with zipfile.ZipFile(io.BytesIO(bytes(data))) as archive:
            names = set(archive.namelist())
    except (zipfile.BadZipFile, OSError):
        return False
    return "xl/workbook.xml" in names


def read_book(
    source: bytes | str | Path,
    *,
    max_bytes: int | None = None,
    max_rows: int | None = None,
    max_cells: int | None = None,
) -> Book:
    """Читает книгу из байтов или из файла.

    Три предела, и все три про чужой файл. `max_bytes` проверяется до
    разбора: слишком большой файл не должен попадать в openpyxl вообще.
    `max_rows` и `max_cells` ограничивают уже разжатое содержимое, потому
    что маленький архив разворачивается в лист на миллионы строк. Без
    значения берутся `MAX_ROWS` и `MAX_CELLS`.
    """
    data = _as_bytes(source)
    if max_bytes is not None and len(data) > int(max_bytes):
        raise TooLargeError(len(data), int(max_bytes))
    if not looks_like_xlsx(data):
        raise NotXlsxError("это не файл xlsx")

    budget = _Budget(
        MAX_ROWS if max_rows is None else max_rows,
        MAX_CELLS if max_cells is None else max_cells,
    )
    try:
        book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as error:  # noqa: BLE001 - селлеру один текст, владельцу причина
        _tell_owner(error)
        raise NotXlsxError("это не файл xlsx") from error

    try:
        return Book(
            tuple(_read_sheet(worksheet, budget) for worksheet in book.worksheets)
        )
    except XlsxError:
        raise
    except Exception as error:  # noqa: BLE001 - то же самое, но на разборе строк
        _tell_owner(error)
        raise NotXlsxError("книга не читается") from error
    finally:
        book.close()


def _tell_owner(error: BaseException) -> None:
    """Настоящая причина уходит в журнал владельца.

    Селлеру показывают «это не xlsx», и для кривого файла это правда. Но
    той же веткой ловится и беда посерьёзнее, вроде нехватки памяти на
    книге-бомбе, а про такое владелец должен узнать.
    """
    audit.log(
        "xlsx",
        None,
        f"Книга не открылась: {type(error).__name__}: {error}",
        level="warning",
    )


def read_sheet(
    source: bytes | str | Path,
    *,
    title: str | None = None,
    max_bytes: int | None = None,
    max_rows: int | None = None,
    max_cells: int | None = None,
) -> SheetData:
    """Один лист: названный или первый, если название не дано."""
    book = read_book(source, max_bytes=max_bytes, max_rows=max_rows, max_cells=max_cells)
    return book[title] if title is not None else book.first


def _as_bytes(source: bytes | str | Path) -> bytes:
    if isinstance(source, (bytes, bytearray)):
        return bytes(source)
    return Path(source).read_bytes()


class _Budget:
    """Сколько строк и ячеек книге ещё разрешено занять в памяти.

    Счёт идёт по всей книге, а не по листу: иначе тысяча листов обошла бы
    предел тысячекратно. Проверка стоит внутри обхода, поэтому огромный лист
    обрывается на пределе, а не после того, как он уже собран.
    """

    def __init__(self, max_rows: int, max_cells: int) -> None:
        self.max_rows = int(max_rows)
        self.max_cells = int(max_cells)
        self.rows = 0
        self.cells = 0

    def take(self, cells: int) -> None:
        self.rows += 1
        self.cells += max(int(cells), 1)
        if self.rows > self.max_rows:
            raise TooManyRowsError(self.rows, self.max_rows)
        if self.cells > self.max_cells:
            raise TooManyCellsError(self.cells, self.max_cells)


def _read_sheet(worksheet: Any, budget: _Budget) -> SheetData:
    rows = worksheet.iter_rows(values_only=True)
    try:
        raw_headers = next(rows)
    except StopIteration:
        return SheetData(str(worksheet.title), (), ())

    headers = [_text(value) for value in raw_headers]
    while headers and not headers[-1]:
        headers.pop()
    width = len(headers)

    found: list[Row] = []
    for number, values in enumerate(rows, start=FIRST_DATA_ROW):
        # Бюджет тратит любая пройденная строка, в том числе пустая: из
        # пустых строк бомба и делается, они сжимаются лучше всего. В Row
        # такая строка по-прежнему не попадает и номера не сдвигает.
        budget.take(len(values))
        if all(value is None or _text(value) == "" for value in values):
            continue
        cells = {
            header: values[index] if index < len(values) else None
            for index, header in enumerate(headers)
            if header
        }
        found.append(Row(number, cells, tuple(values[:width])))

    return SheetData(str(worksheet.title), tuple(headers), tuple(found))


def _text(value: Any) -> str:
    return "" if value is None else str(value).strip()
