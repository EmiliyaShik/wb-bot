"""Общий помощник по xlsx: книги, листы, заголовки.

Про себестоимость тут нет ни слова, и это условие: тем же помощником
пользуются финансовые отчёты, прибыльность артикулов и реестр актов.
"""

from datetime import date
from decimal import Decimal

import pytest

from core import xlsx


def test_write_book_returns_real_xlsx_bytes():
    data = xlsx.write_book(xlsx.Sheet("Лист", ["Имя", "Число"], [["Аня", 7]]))

    assert isinstance(data, bytes)
    assert data[:2] == b"PK"
    assert xlsx.looks_like_xlsx(data)


def test_read_book_returns_headers_and_rows_by_header_name():
    data = xlsx.write_book(
        xlsx.Sheet("Лист", ["Имя", "Число"], [["Аня", 7], ["Боря", 8]])
    )

    sheet = xlsx.read_sheet(data)

    assert sheet.title == "Лист"
    assert sheet.headers == ("Имя", "Число")
    assert [row.cells["Имя"] for row in sheet.rows] == ["Аня", "Боря"]
    assert [row.number for row in sheet.rows] == [2, 3]


def test_book_keeps_several_sheets():
    data = xlsx.write_book(
        [
            xlsx.Sheet("Первый", ["A"], [[1]]),
            xlsx.Sheet("Второй", ["B"], [[2], [3]]),
        ]
    )

    book = xlsx.read_book(data)

    assert book.titles == ("Первый", "Второй")
    assert len(book["Второй"].rows) == 2
    assert book.first.title == "Первый"


def test_read_sheet_by_title_and_unknown_title_is_an_error():
    data = xlsx.write_book(
        [xlsx.Sheet("Первый", ["A"], [[1]]), xlsx.Sheet("Второй", ["B"], [[2]])]
    )

    assert xlsx.read_sheet(data, title="Второй").headers == ("B",)
    with pytest.raises(xlsx.SheetNotFoundError):
        xlsx.read_sheet(data, title="Третий")


def test_column_is_found_by_synonyms_regardless_of_case_and_spaces():
    data = xlsx.write_book(xlsx.Sheet("Лист", ["  nmID ", "Цена, руб"], [[1, 2]]))

    sheet = xlsx.read_sheet(data)

    assert sheet.column("nmid", "код") == "nmID"
    assert sheet.column("название") is None
    assert sheet.missing({"nmid": ("nmid",), "имя": ("имя",)}) == ("имя",)


def test_alien_format_renamed_to_xlsx_is_rejected_not_crashed():
    with pytest.raises(xlsx.NotXlsxError):
        xlsx.read_book(b"%PDF-1.4 and then some bytes")
    assert not xlsx.looks_like_xlsx(b"%PDF-1.4")


def test_zip_without_workbook_inside_is_not_xlsx():
    import io
    import zipfile

    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("readme.txt", "ничего общего с книгой")

    with pytest.raises(xlsx.NotXlsxError):
        xlsx.read_book(buffer.getvalue())


def test_file_above_the_limit_is_rejected_before_parsing():
    data = xlsx.write_book(xlsx.Sheet("Лист", ["A"], [[1]]))

    with pytest.raises(xlsx.TooLargeError):
        xlsx.read_book(data, max_bytes=10)


def test_empty_rows_are_skipped_but_row_numbers_stay_real():
    data = xlsx.write_book(
        xlsx.Sheet("Лист", ["A", "B"], [[1, "x"], [None, None], [3, "z"]])
    )

    sheet = xlsx.read_sheet(data)

    assert [row.number for row in sheet.rows] == [2, 4]


def test_write_book_can_save_to_path_and_read_it_back(tmp_path):
    target = tmp_path / "книга.xlsx"

    xlsx.write_book(xlsx.Sheet("Лист", ["A"], [[42]]), path=target)

    assert target.exists()
    assert xlsx.read_sheet(target).rows[0].cells["A"] == 42


def test_reading_stops_at_the_row_limit_instead_of_swallowing_the_book():
    """Книга на порядок больше предела: важно, что обход оборвался на нём."""
    data = xlsx.write_book(xlsx.Sheet("Лист", ["A"], [[i] for i in range(500)]))

    with pytest.raises(xlsx.TooManyRowsError) as exc:
        xlsx.read_book(data, max_rows=20)

    # Дошли ровно до двадцать первой строки и там остановились, а не собрали
    # все пятьсот, чтобы потом сверить итог.
    assert exc.value.rows == 21
    assert exc.value.limit == 20
    # Предел по умолчанию есть и он конечный: разжатый лист на миллионы
    # строк не должен собираться в память целиком.
    assert 0 < xlsx.MAX_ROWS < 10_000_000


def test_a_sheet_of_empty_rows_spends_the_limit_too():
    """Пустые строки сжимаются лучше всего, и бомбу делают именно из них."""
    data = xlsx.write_book(
        xlsx.Sheet("Лист", ["A"], [[None] for _ in range(500)] + [[1]])
    )

    with pytest.raises(xlsx.TooManyRowsError) as exc:
        xlsx.read_book(data, max_rows=20)

    assert exc.value.rows == 21


def test_reading_stops_at_the_cell_limit_too():
    data = xlsx.write_book(
        xlsx.Sheet("Лист", ["A", "B", "C"], [[1, 2, 3] for _ in range(500)])
    )

    with pytest.raises(xlsx.TooManyCellsError) as exc:
        xlsx.read_book(data, max_cells=30)

    # Одиннадцатая строка перешагнула тридцать ячеек, на ней и встали.
    assert exc.value.cells == 33
    assert 0 < xlsx.MAX_CELLS < 100_000_000


def test_row_limit_counts_the_whole_book_not_one_sheet():
    data = xlsx.write_book(
        [xlsx.Sheet("Первый", ["A"], [[1], [2]]), xlsx.Sheet("Второй", ["A"], [[3]])]
    )

    with pytest.raises(xlsx.TooManyRowsError):
        xlsx.read_book(data, max_rows=2)


# --- деньги в ячейке ---
#
# Excel показывает число так, как его записали: без формата 2290,10 выглядит
# как 2290,1, и две копейки пропадают с экрана. По реестру владелец сверяется
# с банковской выпиской, поэтому копейки в денежной колонке видны всегда.


def money_formats(data: bytes, column: str) -> list[str]:
    """Форматы ячеек одной колонки, кроме заголовка."""
    import io

    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(data))
    try:
        cells = book.worksheets[0][column]
        return [cell.number_format for cell in cells[xlsx.HEADER_ROW:]]
    finally:
        book.close()


def test_a_column_signed_in_rubles_shows_kopecks():
    data = xlsx.write_book(
        xlsx.Sheet("Лист", ["Номер", "Сумма, ₽"], [["С-1", Decimal("2290.10")]])
    )

    assert money_formats(data, "B") == [xlsx.MONEY_FORMAT]
    # И число при этом осталось числом, а не текстом: его складывают в Excel.
    assert xlsx.read_sheet(data).rows[0].cells["Сумма, ₽"] == 2290.1


def test_a_column_without_the_ruble_sign_is_left_as_it_was():
    """Формат ставится по подписи колонки, а не всему числовому подряд."""
    data = xlsx.write_book(
        xlsx.Sheet("Лист", ["Период, мес.", "Сумма, ₽"], [[3, Decimal("990.00")]])
    )

    assert money_formats(data, "A") == ["General"]
    assert money_formats(data, "B") == [xlsx.MONEY_FORMAT]


def test_the_ruble_sign_is_recognized_with_a_trailing_space():
    assert xlsx.is_money_header("Сумма, ₽ ")
    assert xlsx.is_money_header("Себестоимость за единицу, ₽")
    assert not xlsx.is_money_header("Комиссия WB, %")
    assert not xlsx.is_money_header(None)


def test_an_empty_money_cell_keeps_the_format_for_what_the_seller_types():
    """Шаблон себестоимости уходит селлеру пустым, и вписывать он будет деньги."""
    data = xlsx.write_book(xlsx.Sheet("Лист", ["Себестоимость, ₽"], [[None]]))

    assert money_formats(data, "A") == [xlsx.MONEY_FORMAT]


def test_kopecks_survive_the_trip_through_the_cell():
    """openpyxl печатает число через %.16g, то есть Decimal идёт в файл
    через float. Для денег этого хватает с большим запасом, и обратно
    читается ровно та же сумма."""
    data = xlsx.write_book(
        xlsx.Sheet("Лист", ["Сумма, ₽"], [[Decimal(cents) / 100] for cents in (229010, 1, 99999999)])
    )

    got = [row.cells["Сумма, ₽"] for row in xlsx.read_sheet(data).rows]

    assert [Decimal(str(value)) for value in got] == [
        Decimal("2290.10"),
        Decimal("0.01"),
        Decimal("999999.99"),
    ]


def test_unreadable_workbook_tells_the_owner_the_real_reason(monkeypatch):
    """Селлер видит «это не xlsx», владелец в журнале - настоящую причину."""
    written = []
    monkeypatch.setattr(
        xlsx.audit, "log", lambda kind, client_id, message, **kw: written.append(message)
    )

    def explode(*args, **kwargs):
        raise ValueError("внутренняя беда разбора")

    monkeypatch.setattr(xlsx, "load_workbook", explode)
    data = xlsx.write_book(xlsx.Sheet("Лист", ["A"], [[1]]))

    with pytest.raises(xlsx.NotXlsxError):
        xlsx.read_book(data)

    assert any("внутренняя беда разбора" in message for message in written)


# --- колонка, которой нечем заполниться --------------------------------------
#
# Правило одно на обе книги проекта: и финансовую раскладку, и прибыльность
# артикулов. Живёт оно здесь, потому что здесь живут книги: `core.excel` про
# состав финансовых листов, `agents.profit` про прибыль, и ни один из них не
# годится в общее место для второго. Копия такого правила в двух файлах
# разошлась бы молча, и за одну и ту же неделю две книги показали бы разные
# колонки.


def test_a_statement_that_never_reached_a_row_is_hidden_and_a_true_zero_is_not():
    # Хранение есть, но всё пришло строками без артикула: колонке нечем
    # заполниться. Приёмки не было вовсе: там ноль настоящий.
    assert xlsx.only_faceless(Decimal("0"), Decimal("600")) is True
    assert xlsx.only_faceless(Decimal("0"), Decimal("0")) is False
    assert xlsx.only_faceless(Decimal("5"), Decimal("600")) is False
    # Число может приехать и строкой из базы, и None из пустой клетки.
    assert xlsx.only_faceless(None, "600") is True
    assert xlsx.only_faceless("0.00", None) is False


def test_hidden_keys_and_the_columns_that_go_with_them():
    hidden = xlsx.faceless_keys(
        {
            "storage": (Decimal("0"), Decimal("600")),
            "logistics": (Decimal("75"), Decimal("0")),
            "acceptance": (Decimal("0"), Decimal("0")),
        }
    )
    columns = (
        ("Артикул WB", None),
        ("Логистика, ₽", "logistics"),
        ("Хранение, ₽", "storage"),
        ("Приёмка, ₽", "acceptance"),
    )

    assert hidden == frozenset({"storage"})
    assert xlsx.visible_headers(columns, hidden) == (
        "Артикул WB",
        "Логистика, ₽",
        "Приёмка, ₽",
    )
    # Значения уходят из строки ровно те же и в том же порядке.
    assert xlsx.visible_row([111, 75, 600, 0], columns, hidden) == [111, 75, 0]


def finance_report():
    """Финансовый отчёт в памяти: хранение пришло строкой без артикула."""
    from agents.finance import Amounts, Article, FinanceReport, Week

    sold = Amounts(revenue=Decimal("1000"), logistics=Decimal("50"), sales_count=1)
    faceless = Amounts(storage=Decimal("600"))
    return FinanceReport(
        client_id=1,
        period="week",
        date_from=date(2026, 9, 7),
        date_to=date(2026, 9, 13),
        weeks=(
            Week(
                report_id=1,
                date_from="2026-09-07",
                date_to="2026-09-13",
                amounts=sold + faceless,
            ),
        ),
        months=(),
        articles=(
            Article(nm_id=111, vendor_code="ART-1", quantity=1, amounts=sold),
            Article(nm_id=None, amounts=faceless),
        ),
    )


def profit_report():
    """Отчёт о прибыльности в памяти: та же неделя, та же дыра."""
    from agents.profit import ArticleProfit, CostItem, ProfitReport

    return ProfitReport(
        client_id=1,
        period="week",
        date_from=date(2026, 9, 7),
        date_to=date(2026, 9, 13),
        articles=(
            ArticleProfit(nm_id=111, vendor_code="ART-1", logistics=Decimal("50")),
        ),
        unallocated=Decimal("600"),
        costs_breakdown=(
            CostItem(key="logistics", title="Логистика", by_article=Decimal("50")),
            CostItem(key="storage", title="Хранение", faceless=Decimal("600")),
        ),
    )


def book_headers():
    """Заголовки листа артикулов в обеих книгах проекта."""
    from agents import profit
    from core import excel

    return (
        xlsx.read_book(excel.finance_book(finance_report()))["По артикулам"].headers,
        profit.article_headers(profit_report()),
    )


def test_both_books_hide_the_column_by_one_and_the_same_rule(monkeypatch):
    """Сломайте правило здесь, и изменятся сразу обе книги.

    Это и есть доказательство, что копии у него нет: финансовая раскладка и
    прибыльность спрашивают одну функцию, а не помнят каждая своё.
    """
    for headers in book_headers():
        assert "Хранение, ₽" not in headers
        assert "Логистика, ₽" in headers

    # Правило сломано ровно в одном месте.
    monkeypatch.setattr(xlsx, "only_faceless", lambda by_row, faceless: False)

    for headers in book_headers():
        assert "Хранение, ₽" in headers
